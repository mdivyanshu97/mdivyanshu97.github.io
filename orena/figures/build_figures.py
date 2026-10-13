#!/usr/bin/env python3
"""Build the ORena portfolio figures and slides without patient imagery.

Run from any directory with matplotlib and python-pptx installed:
    python orena/figures/build_figures.py
"""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.image as mpimg
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
INK = "#142b38"
MUTED = "#53636a"
BLUE = "#235b76"
ORANGE = "#b8512e"
LINE = "#d9ddd9"
SOFT_BLUE = "#eaf0f2"
SOFT_ORANGE = "#f6ece5"

plt.rcParams.update(
    {
        "font.family": "DejaVu Sans",
        "font.size": 13,
        "svg.fonttype": "path",
        "pdf.fonttype": 42,
        "axes.unicode_minus": False,
    }
)


def canvas(height=460):
    fig = plt.figure(figsize=(12, height / 100), dpi=100, facecolor="white")
    ax = fig.add_axes((0, 0, 1, 1))
    ax.set(xlim=(0, 1200), ylim=(height, 0))
    ax.axis("off")
    return fig, ax


def label(ax, x, y, text, size=14, color=INK, weight="normal", max_width=None, **kwargs):
    artist = ax.text(
        x,
        y,
        text,
        va="top",
        color=color,
        fontsize=size,
        fontweight=weight,
        linespacing=1.5,
        **kwargs,
    )
    if max_width is not None:
        renderer = ax.figure.canvas.get_renderer()
        while size > 9:
            bounds = artist.get_window_extent(renderer).transformed(ax.transData.inverted())
            if bounds.width <= max_width:
                break
            size -= 0.5
            artist.set_fontsize(size)
    return artist


def card(ax, x, y, width, height, fill="white", edge=LINE):
    ax.add_patch(
        FancyBboxPatch(
            (x, y),
            width,
            height,
            boxstyle="round,pad=0,rounding_size=8",
            facecolor=fill,
            edgecolor=edge,
            linewidth=1,
        )
    )


def arrow(ax, start, end, color=BLUE, style="-|>", width=1.5):
    ax.add_patch(
        FancyArrowPatch(
            start,
            end,
            arrowstyle=style,
            mutation_scale=13,
            lw=width,
            color=color,
        )
    )


def save(fig, name):
    for ext in ("svg", "pdf", "png"):
        output = ROOT / f"{name}.{ext}"
        fig.savefig(
            output,
            dpi=180 if ext == "png" else 100,
            facecolor="white",
            metadata={"Creator": "Divyanshu Mishra"} if ext == "pdf" else None,
        )
        if ext == "svg":
            output.write_text(
                "\n".join(line.rstrip() for line in output.read_text().splitlines()) + "\n"
            )
    plt.close(fig)


def training():
    fig, ax = canvas()
    label(ax, 26, 20, "SUPERVISED MULTIMODAL POST-TRAINING", 12, BLUE, "bold")
    card(ax, 26, 79, 245, 229)
    label(ax, 46, 99, "01 / Surgical warm-up", 13, BLUE, "bold", max_width=205)
    label(ax, 46, 140, "Qwen3-VL-4B", 22, INK, "bold", max_width=205)
    label(ax, 46, 188, "SSG-VQA questions\n+ CholecT45 frames", 16, max_width=205)
    label(ax, 46, 257, "Learn surgical scene literacy", 11, MUTED, max_width=205)
    arrow(ax, (278, 193), (323, 193))

    card(ax, 331, 79, 318, 229, SOFT_BLUE)
    label(ax, 352, 99, "02 / Joint task fine-tuning", 13, BLUE, "bold", max_width=276)
    label(ax, 352, 140, "FRAME · SEGMENT\n· PROCEDURE", 20, INK, "bold", max_width=276)
    label(
        ax,
        352,
        211,
        "Image-conditioned chat examples\nLoss on answer tokens only",
        14,
        max_width=276,
    )
    label(ax, 352, 267, "Continue the surgical adapters", 11, MUTED)
    arrow(ax, (654, 193), (719, 193))
    ax.plot([715, 715], [138, 248], color=BLUE, lw=1.5)
    arrow(ax, (715, 138), (756, 138))
    arrow(ax, (715, 248), (756, 248))

    card(ax, 763, 79, 407, 108)
    label(ax, 782, 96, "03 / SEGMENT", 12, BLUE, "bold")
    label(ax, 782, 125, "Language + vision + merger LoRA", 16, INK, "bold", max_width=368)
    label(ax, 782, 157, "One merged checkpoint", 12, MUTED)
    card(ax, 763, 203, 407, 105, SOFT_ORANGE)
    label(ax, 782, 219, "03 / PROCEDURE", 12, ORANGE, "bold")
    label(ax, 782, 246, "Independent W64 and NW2 adapters", 15, INK, "bold", max_width=368)
    label(ax, 782, 275, "Language LoRA · visual backbone frozen", 12, MUTED, max_width=368)

    ax.plot([26, 1170], [344, 344], color=LINE, lw=1)
    label(ax, 26, 365, "LoRA r16 / α32", 15, INK, "bold")
    label(ax, 26, 399, "Segment visual mergers use r64", 11, MUTED)
    label(ax, 444, 365, "Matched visual preprocessing", 15, INK, "bold", max_width=380)
    label(ax, 444, 399, "Chronological frames · absolute-time clock", 11, MUTED, max_width=380)
    label(ax, 856, 365, "Merge for local inference", 15, INK, "bold", max_width=314)
    label(ax, 856, 399, "Independent bfloat16 model artifacts", 11, MUTED, max_width=314)
    save(fig, "training")


def film(ax, x, y, width, count=32, color=SOFT_BLUE, height=31):
    gap = width / count
    for index in range(count):
        ax.add_patch(
            Rectangle(
                (x + index * gap, y),
                max(1, gap - 3),
                height,
                facecolor=color,
                edgecolor=LINE,
                linewidth=0.5,
            )
        )


def search():
    fig, ax = canvas()
    label(ax, 26, 18, "COARSE-TO-FINE TEMPORAL INFERENCE", 12, BLUE, "bold")
    label(ax, 26, 69, "Whole procedure", 19, INK, "bold")
    label(ax, 26, 103, "W64 · 128 frames", 13, MUTED)
    film(ax, 325, 78, 829, count=48)
    card(ax, 702, 72, 270, 43, SOFT_ORANGE, ORANGE)
    label(ax, 708, 122, "Initial pointer p₀", 12, ORANGE)

    ax.plot([702, 522], [115, 212], color=ORANGE, lw=1.2)
    ax.plot([972, 1050], [115, 212], color=ORANGE, lw=1.2)
    label(ax, 26, 193, "20-minute window", 19, INK, "bold")
    label(ax, 26, 228, "NW2 · 128 frames", 13, MUTED)
    film(ax, 522, 213, 528, count=37)
    card(ax, 819, 207, 92, 43, SOFT_ORANGE, ORANGE)
    label(ax, 770, 258, "Pointer p₁", 12, ORANGE)

    ax.plot([819, 667], [250, 343], color=ORANGE, lw=1.2)
    ax.plot([911, 1037], [250, 343], color=ORANGE, lw=1.2)
    label(ax, 26, 326, "100-second window", 19, INK, "bold")
    label(ax, 26, 361, "NW2 · 64 frames", 13, MUTED)
    film(ax, 667, 344, 370, count=32, color=SOFT_ORANGE)
    ax.plot([891, 891], [333, 383], color=ORANGE, lw=2)
    label(ax, 849, 391, "Final hh:mm:ss", 13, ORANGE, "bold")
    label(
        ax,
        26,
        420,
        "One model resident at a time · absolute video clock across all passes · schematic windows",
        11,
        MUTED,
    )
    save(fig, "search")


def results():
    data = json.loads((ROOT / "results.json").read_text())
    fig, ax = canvas(480)
    label(ax, 26, 17, "LOCAL OFFLINE ABLATIONS", 12, BLUE, "bold")
    x0, x1 = 548, 1064
    for tick in (0, 25, 50, 75, 100):
        x = x0 + tick / 100 * (x1 - x0)
        ax.plot([x, x], [76, 373], color=LINE, lw=0.8, zorder=0)
        label(ax, x, 378, str(tick), 11, MUTED, ha="center")
    label(
        ax,
        (x0 + x1) / 2,
        407,
        "Accuracy on the indicated question subset (%)",
        11,
        MUTED,
        ha="center",
    )
    for i, study in enumerate(data["studies"]):
        y = 92 + 105 * i
        label(ax, 26, y - 12, study["track"].upper(), 11, BLUE, "bold")
        label(ax, 26, y + 13, study["label"], 16, INK, "bold")
        label(ax, 26, y + 44, f"n = {study['n']:,} · {study['subset']}", 11, MUTED)
        before = x0 + study["baseline"] / 100 * (x1 - x0)
        after = x0 + study["improved"] / 100 * (x1 - x0)
        arrow(ax, (before, y + 30), (after, y + 30), ORANGE, width=2.3)
        ax.scatter([before, after], [y + 30, y + 30], s=[74, 92], c=[BLUE, ORANGE], zorder=3)
        label(ax, before - 7, y - 4, f"{study['baseline']:.1f}", 14, BLUE, "bold", ha="right")
        label(ax, after + 7, y - 4, f"{study['improved']:.1f}", 14, ORANGE, "bold")
        label(
            ax,
            1170,
            y + 14,
            f"+{study['improved'] - study['baseline']:.1f} pp",
            15,
            ORANGE,
            "bold",
            ha="right",
        )
    label(
        ax,
        26,
        447,
        "Separate question sets · Procedure search study uses one checkpoint, before the W64/NW2 split",
        11,
        MUTED,
    )
    save(fig, "results")


def slides():
    presentation = Presentation()
    presentation.slide_width = Inches(16)
    presentation.slide_height = Inches(9)
    content = [
        (
            "Surgical VLM post-training",
            "training",
            "A solo project by Divyanshu Mishra · Incision Impossible · ORena SAVE FOCUS 2026",
            "Two-stage supervised adaptation: surgical scene literacy followed by joint FOCUS fine-tuning. "
            "The Segment and Procedure checkpoints differ in their trainable visual components. "
            "Training code and full manifests are not included in the public inference release.",
        ),
        (
            "Answer, then zoom in",
            "search",
            "W64 → NW2 · local inference · one model resident at a time",
            "Procedure uses W64 for the full-video pointer and NW2 for both zoom passes. "
            "The three stages sample 128, 128 and 64 frames. Windows are 1200 and 100 seconds. "
            "The diagram is schematic and contains no patient imagery.",
        ),
        (
            "Evaluate the evidence policy",
            "results",
            "Local offline studies · separate subsets · gains are not additive",
            "Segment refinement: 50.4 to 55.0 percent on 1577 single-timestamp questions, earlier checkpoint. "
            "Segment question windows: 72.7 to 79.1 percent on 422 questions with stated times. "
            "Procedure search: 6.8 to 35.1 percent on 746 single-timestamp questions, one checkpoint, "
            "before the W64/NW2 split. These are not official challenge scores.",
        ),
    ]
    with PdfPages(
        ROOT.parent / "spotlight.pdf",
        metadata={
            "Title": "ORena surgical video QA",
            "Author": "Divyanshu Mishra",
        },
    ) as pdf:
        for i, (title, image, subtitle, notes) in enumerate(content):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(14), Inches(0.8))
            run = title_box.text_frame.paragraphs[0].add_run()
            run.text = title
            run.font.name = "Arial"
            run.font.size = Pt(32)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string("142B38")
            slide.shapes.add_picture(
                str(ROOT / f"{image}.png"), Inches(0.6), Inches(1.8), width=Inches(14.8)
            )
            footer = slide.shapes.add_textbox(Inches(0.7), Inches(8.1), Inches(14), Inches(0.5))
            paragraph = footer.text_frame.paragraphs[0]
            paragraph.text = subtitle
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(14)
            paragraph.font.color.rgb = RGBColor.from_string("53636A")
            slide.notes_slide.notes_text_frame.text = notes
            fig = plt.figure(figsize=(16, 9), facecolor="white")
            fig.text(0.045, 0.86, title, fontsize=28, fontweight="bold", color=INK)
            ax = fig.add_axes((0.038, 0.16, 0.925, 0.65))
            ax.imshow(mpimg.imread(ROOT / f"{image}.png"))
            ax.axis("off")
            fig.text(0.045, 0.065, subtitle, fontsize=12, color=MUTED)
            pdf.savefig(fig)
            plt.close(fig)
    presentation.save(ROOT.parent / "spotlight.pptx")


def main():
    training()
    search()
    results()
    slides()
    print("Built three SVG/PDF/PNG figures and the PDF/PowerPoint project slides.")


if __name__ == "__main__":
    main()
